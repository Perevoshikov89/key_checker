# -*- coding: utf-8 -*-
"""NBKI key-checker 1.0: DB2 BEFORE/AFTER baseline comparison.
Requires: Python 3.10+, ibm_db. No loader-log parsing is performed.
"""
from __future__ import annotations
import csv, json, os, re, html, hashlib, threading, traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import xml.etree.ElementTree as ET

APP_TITLE = "NBKI Key Checker 1.0 — DB2 baseline"
DEFAULT_HOST = "10.230.227.100"
DEFAULT_PORT = "2668"
DEFAULT_DATABASE = "cprosd22"
DEFAULT_SCHEMA = "INDIC"
DEFAULT_USER = "yperevos"
DEFAULT_PASSWORD = ""
DEFAULT_FID_COLUMN = "FID"
IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_$#]*$")


def now_stamp():
    return datetime.now().strftime("%Y-%m-%d_%H-%M-%S")


def safe_ident(value: str, label: str) -> str:
    value = (value or "").strip().upper()
    if not IDENT_RE.fullmatch(value):
        raise ValueError(f"Некорректное имя {label}: {value!r}")
    return value


def read_mapping(path: str):
    """Read mapping from Excel (.xlsx) or CSV/TSV, with Russian or English headers."""
    if not path:
        raise ValueError("Выберите Excel, CSV или TSV-файл конфигурации показателей.")

    source = Path(path)
    suffix = source.suffix.lower()
    rows = []
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Для чтения Excel .xlsx установите openpyxl: python -m pip install openpyxl")
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("В Excel-файле нет заголовков.")
            headers = [str(v).strip() if v is not None else "" for v in headers]
            for values_row in values:
                rows.append({headers[i]: (values_row[i] if i < len(values_row) else None)
                             for i in range(len(headers)) if headers[i]})
            workbook.close()
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Не удалось прочитать Excel-файл: {e}")
        fieldnames = headers
    else:
        if suffix not in (".csv", ".tsv", ".txt"):
            raise ValueError("Поддерживаются файлы .xlsx, .csv, .tsv и .txt.")
        raw = source.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1251", "utf-8"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise ValueError("Не удалось определить кодировку файла конфига.")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in sample else csv.excel
        reader = csv.DictReader(text.splitlines(), dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("В файле конфигурации нет заголовков.")
        fieldnames = reader.fieldnames
        rows = list(reader)

    norm = {re.sub(r"[^a-zа-я0-9]+", "", str(k).lower()): k for k in fieldnames if k}
    def find(*keys):
        for key in keys:
            k = re.sub(r"[^a-zа-я0-9]+", "", key.lower())
            if k in norm: return norm[k]
        return None
    c_indicator = find("№ п-ля", "номер поля", "indicator", "indicator number", "поле номер")
    c_name = find("Имя показателя", "name", "indicator name")
    c_table = find("Таблица", "table", "table name")
    c_field = find("Поле", "field", "column", "column name")
    if not c_table or not c_field:
        raise ValueError("В конфиге должны быть столбцы 'Таблица' и 'Поле'.")

    def cell(row, column):
        value = row.get(column, "") if column else ""
        return "" if value is None else str(value).strip()

    mapping = []
    seen = set()
    for row in rows:
        table_raw, field_raw = cell(row, c_table), cell(row, c_field)
        if not table_raw or not field_raw:
            continue
        table = safe_ident(table_raw, "таблицы")
        field = safe_ident(field_raw, "поля")
        key = (table, field)
        if key in seen: continue
        seen.add(key)
        mapping.append({"indicator": cell(row, c_indicator), "name": cell(row, c_name),
                        "table": table, "field": field})
    if not mapping:
        raise ValueError("В конфиге не найдено ни одной пары таблица/поле.")
    grouped = {}
    for item in mapping:
        grouped.setdefault(item["table"], set()).add(item["field"])
    return mapping, {t: sorted(fields) for t, fields in grouped.items()}


def get_event_comments(xml_dir: str):
    comments = []
    if not xml_dir:
        return comments
    root = Path(xml_dir)
    if not root.exists():
        raise ValueError("Папка XML не найдена.")
    for path in sorted(root.rglob("*.xml")):
        try:
            tree = ET.parse(path)
            found = []
            for elem in tree.iter():
                for key, value in elem.attrib.items():
                    if key.split("}")[-1].lower() == "eventcomment":
                        found.append(value)
                # Also handle an eventComment element's text if XML uses an element rather than an attribute.
                if elem.tag.split("}")[-1].lower() == "eventcomment" and elem.text:
                    found.append(elem.text)
            comments.append({"file": path.name, "comments": found, "error": ""})
        except Exception as e:
            comments.append({"file": path.name, "comments": [], "error": str(e)})
    return comments


class DB2Client:
    def __init__(self, host, port, database, user, password, schema, cli_path=""):
        self.host, self.port, self.database = host.strip(), port.strip(), database.strip()
        self.user, self.password = user.strip(), password
        self.schema = safe_ident(schema, "схемы")
        self.cli_path = cli_path.strip()

    def connect(self):
        try:
            import ibm_db
        except ImportError:
            raise RuntimeError("Не найден модуль ibm_db. Установите его в используемый Python: python -m pip install ibm_db")
        if self.cli_path:
            os.environ["IBM_DB_HOME"] = self.cli_path
            os.environ["DB2_CLI_DRIVER_INSTALL_PATH"] = self.cli_path
        parts = [f"DATABASE={self.database}", f"HOSTNAME={self.host}", f"PORT={self.port}", "PROTOCOL=TCPIP", f"UID={self.user}", f"PWD={self.password}"]
        conn = ibm_db.connect(";".join(parts) + ";", "", "")
        return ibm_db, conn

    def snapshot(self, fid: str, grouped_fields: dict, fid_column: str):
        fid_column = safe_ident(fid_column, "поля FID")
        ibm_db, conn = self.connect()
        result = {"created_at": datetime.now().isoformat(timespec="seconds"),
                  "host": self.host, "database": self.database, "schema": self.schema,
                  "fid": fid, "fid_column": fid_column, "tables": {}}
        try:
            for table, fields in grouped_fields.items():
                # Verify table and requested fields exist in the configured schema.
                stmt = ibm_db.exec_immediate(conn,
                    "SELECT COLNAME FROM SYSCAT.COLUMNS WHERE TABSCHEMA='" + self.schema + "' AND TABNAME='" + table + "'")
                available = set()
                rec = ibm_db.fetch_assoc(stmt)
                while rec:
                    available.add(str(rec["COLNAME"]).upper())
                    rec = ibm_db.fetch_assoc(stmt)
                missing = [x for x in [fid_column] + fields if x not in available]
                if missing:
                    raise RuntimeError(f"{self.schema}.{table}: отсутствуют поля {', '.join(missing)}. Проверьте схему и колонку FID.")
                cols = sorted(set(fields + [fid_column]))
                select_cols = ", ".join('"' + c + '"' for c in cols)
                sql = f'SELECT {select_cols} FROM "{self.schema}"."{table}" WHERE "{fid_column}" = ?'
                stmt = ibm_db.prepare(conn, sql)
                ibm_db.bind_param(stmt, 1, fid)
                ibm_db.execute(stmt)
                rows = []
                row = ibm_db.fetch_assoc(stmt)
                while row:
                    clean = {}
                    for k, v in row.items():
                        if hasattr(v, "isoformat"):
                            v = v.isoformat()
                        elif isinstance(v, bytes):
                            v = v.decode("utf-8", errors="replace")
                        clean[str(k).upper()] = v
                    rows.append(clean)
                    row = ibm_db.fetch_assoc(stmt)
                result["tables"][table] = {"fields": cols, "rows": rows, "count": len(rows)}
            return result
        finally:
            ibm_db.close(conn)


def row_key(row):
    return json.dumps(row, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))


def compare_snapshots(before, after):
    changes = []
    all_tables = sorted(set(before.get("tables", {})) | set(after.get("tables", {})))
    for table in all_tables:
        b = before.get("tables", {}).get(table, {}).get("rows", [])
        a = after.get("tables", {}).get(table, {}).get("rows", [])
        bc, ac = Counter(map(row_key, b)), Counter(map(row_key, a))
        for encoded, n in (ac - bc).items():
            for _ in range(n): changes.append({"table": table, "change": "Добавлена запись", "row": json.loads(encoded)})
        for encoded, n in (bc - ac).items():
            for _ in range(n): changes.append({"table": table, "change": "Удалена запись", "row": json.loads(encoded)})
        b_count, a_count = len(b), len(a)
        if b_count != a_count:
            # Table summary appears separately from row-level details.
            pass
    return changes


def html_report(before, after, changes, comments, out_path):
    esc = lambda x: html.escape(str(x if x is not None else ""))
    table_rows = []
    for table in sorted(set(before.get("tables", {})) | set(after.get("tables", {}))):
        bc = before.get("tables", {}).get(table, {}).get("count", 0)
        ac = after.get("tables", {}).get(table, {}).get("count", 0)
        delta = ac - bc
        table_rows.append(f"<tr><td>{esc(table)}</td><td>{bc}</td><td>{ac}</td><td>{delta:+}</td></tr>")
    change_rows = []
    for c in changes:
        values = "<br>".join(f"<b>{esc(k)}</b>: {esc(v)}" for k, v in c["row"].items())
        change_rows.append(f"<tr><td>{esc(c['table'])}</td><td>{esc(c['change'])}</td><td>{values}</td></tr>")
    comment_rows = []
    for c in comments:
        val = "<br><br>".join(esc(x) for x in c["comments"]) or "<i>eventComment не найден</i>"
        if c["error"]: val += "<br><b>Ошибка чтения XML:</b> " + esc(c["error"])
        comment_rows.append(f"<tr><td>{esc(c['file'])}</td><td>{val}</td></tr>")
    html_doc = f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>DB2 Case Checker report</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;margin:28px;color:#222}}h1,h2{{color:#222}}table{{border-collapse:collapse;width:100%;margin:12px 0 28px}}th,td{{border:1px solid #bbb;padding:8px;text-align:left;vertical-align:top;overflow-wrap:anywhere}}th{{background:#eee}}.meta{{background:#f6f6f6;padding:12px;border-radius:6px}}.empty{{color:#666}}</style></head><body>
<h1>DB2 Case Checker — отчёт сравнения baseline</h1><div class="meta"><b>FID:</b> {esc(before.get('fid'))}<br><b>Схема:</b> {esc(before.get('schema'))}<br><b>BEFORE:</b> {esc(before.get('created_at'))}<br><b>AFTER:</b> {esc(after.get('created_at'))}<br><b>Таблиц:</b> {len(table_rows)}<br><b>Изменений:</b> {len(changes)}</div>
<h2>Сводка по таблицам</h2><table><thead><tr><th>Таблица</th><th>BEFORE</th><th>AFTER</th><th>Изменение</th></tr></thead><tbody>{''.join(table_rows)}</tbody></table>
<h2>Обнаруженные изменения</h2><table><thead><tr><th>Таблица</th><th>Тип изменения</th><th>Значения полей записи</th></tr></thead><tbody>{''.join(change_rows) if change_rows else '<tr><td colspan="3" class="empty">Изменений не обнаружено</td></tr>'}</tbody></table>
<h2>XML-файлы и eventComment</h2><table><thead><tr><th>Файл</th><th>Полный eventComment</th></tr></thead><tbody>{''.join(comment_rows) if comment_rows else '<tr><td colspan="2" class="empty">XML-файлы не выбраны</td></tr>'}</tbody></table>
<p class="empty">Отчёт показывает фактическую разницу BEFORE/AFTER. eventComment приведён только для пояснения и не интерпретируется программой.</p></body></html>'''
    Path(out_path).write_text(html_doc, encoding="utf-8")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE); self.geometry("980x760"); self.minsize(900, 650)
        self.mapping_path = tk.StringVar(); self.xml_dir = tk.StringVar(); self.output_dir = tk.StringVar(value=str(Path.cwd()))
        self.host = tk.StringVar(value=DEFAULT_HOST); self.port = tk.StringVar(value=DEFAULT_PORT)
        self.database = tk.StringVar(value=DEFAULT_DATABASE); self.user = tk.StringVar(value=DEFAULT_USER); self.password = tk.StringVar(value=DEFAULT_PASSWORD)
        self.schema = tk.StringVar(value=DEFAULT_SCHEMA); self.fid_column = tk.StringVar(value=DEFAULT_FID_COLUMN)
        self.fid = tk.StringVar(); self.cli_path = tk.StringVar()
        self.before_path = tk.StringVar(); self.status = tk.StringVar(value="Готов к работе.")
        self._build()

    def _build(self):
        root = ttk.Frame(self, padding=12); root.pack(fill="both", expand=True)
        db = ttk.LabelFrame(root, text="Подключение DB2", padding=10); db.pack(fill="x")
        def entry(parent, label, var, row, col, width=22, show=None):
            ttk.Label(parent, text=label).grid(row=row, column=col, sticky="w", padx=4, pady=4)
            ttk.Entry(parent, textvariable=var, width=width, show=show).grid(row=row, column=col+1, sticky="ew", padx=4, pady=4)
        entry(db,"Host",self.host,0,0); entry(db,"Port",self.port,0,2)
        entry(db,"Database",self.database,1,0); entry(db,"Schema",self.schema,1,2)
        entry(db,"User",self.user,2,0); entry(db,"Password",self.password,2,2,show="*")
        entry(db,"Колонка FID",self.fid_column,3,0); entry(db,"FID субъекта",self.fid,3,2)
        ttk.Label(db,text="DB2 CLI Driver (необязательно)").grid(row=4,column=0,sticky="w",padx=4,pady=4)
        ttk.Entry(db,textvariable=self.cli_path).grid(row=4,column=1,columnspan=2,sticky="ew",padx=4,pady=4)
        ttk.Button(db,text="Обзор…",command=self.pick_cli).grid(row=4,column=3,padx=4)
        for c in range(4): db.columnconfigure(c,weight=1)
        files = ttk.LabelFrame(root,text="Входные данные",padding=10); files.pack(fill="x",pady=10)
        self._path_row(files,"Конфиг показателей (Excel/CSV/TSV)",self.mapping_path,self.pick_mapping,0)
        self._path_row(files,"Папка с XML (для eventComment)",self.xml_dir,self.pick_xml,1)
        self._path_row(files,"Папка для baseline и отчёта",self.output_dir,self.pick_output,2)
        actions = ttk.LabelFrame(root,text="Действия",padding=10); actions.pack(fill="x")
        ttk.Button(actions,text="Проверить подключение",command=self.run_test).pack(side="left",padx=4)
        ttk.Button(actions,text="1. Снять BEFORE",command=self.run_before).pack(side="left",padx=4)
        ttk.Button(actions,text="2. Снять AFTER и сравнить",command=self.run_after).pack(side="left",padx=4)
        ttk.Label(actions,textvariable=self.before_path).pack(side="left",padx=8)
        logbox = ttk.LabelFrame(root,text="Журнал работы",padding=8); logbox.pack(fill="both",expand=True,pady=10)
        self.log = tk.Text(logbox,height=12,wrap="word",state="disabled"); self.log.pack(fill="both",expand=True)
        ttk.Label(root,textvariable=self.status).pack(anchor="w")
        self._log("Кей-чекер не читает лог загрузчика. eventComment используется только в отчёте.")

    def _path_row(self, parent, label, var, command, row):
        ttk.Label(parent,text=label).grid(row=row,column=0,sticky="w",padx=4,pady=4)
        ttk.Entry(parent,textvariable=var).grid(row=row,column=1,sticky="ew",padx=4,pady=4)
        ttk.Button(parent,text="Обзор…",command=command).grid(row=row,column=2,padx=4,pady=4)
        parent.columnconfigure(1,weight=1)
    def pick_mapping(self): self.mapping_path.set(filedialog.askopenfilename(filetypes=[("Excel и CSV/TSV","*.xlsx *.csv *.tsv *.txt"),("Excel (*.xlsx)","*.xlsx"),("CSV (*.csv)","*.csv"),("TSV/TXT","*.tsv *.txt"),("Все файлы","*.*")]))
    def pick_xml(self): self.xml_dir.set(filedialog.askdirectory())
    def pick_output(self): self.output_dir.set(filedialog.askdirectory())
    def pick_cli(self): self.cli_path.set(filedialog.askdirectory())
    def _log(self,msg):
        self.log.configure(state="normal"); self.log.insert("end",f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n"); self.log.see("end"); self.log.configure(state="disabled")
    def _client(self):
        if not self.database.get().strip(): raise ValueError("Укажите имя базы DB2. Хост/алиас и имя БД — разные параметры.")
        if not self.user.get().strip(): raise ValueError("Укажите пользователя DB2.")
        return DB2Client(self.host.get(),self.port.get(),self.database.get(),self.user.get(),self.password.get(),self.schema.get(),self.cli_path.get())
    def _inputs(self):
        if not self.fid.get().strip(): raise ValueError("Укажите FID субъекта.")
        mapping, grouped = read_mapping(self.mapping_path.get())
        return self._client(), mapping, grouped, self.fid.get().strip()
    def _thread(self, fn):
        def work():
            try: fn()
            except Exception as e:
                self.after(0,lambda e=e: (self._log("ОШИБКА: "+str(e)), self._log(traceback.format_exc()), self.status.set("Операция завершилась ошибкой."), messagebox.showerror("Ошибка",str(e))))
        threading.Thread(target=work,daemon=True).start()
    def run_test(self):
        def task():
            client=self._client(); ibm_db,conn=client.connect(); ibm_db.close(conn)
            self.after(0,lambda:(self._log("Подключение DB2 успешно."),self.status.set("Подключение DB2 успешно.")))
        self._thread(task)
    def run_before(self):
        def task():
            client,mapping,grouped,fid=self._inputs()
            self.after(0,lambda:self._log(f"Снимаю BEFORE по FID={fid}; таблиц: {len(grouped)}"))
            snap=client.snapshot(fid,grouped,self.fid_column.get())
            path=Path(self.output_dir.get())/f"case_baseline_BEFORE_{now_stamp()}.json"
            path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(snap,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
            self.before_path.set(str(path))
            self.after(0,lambda:(self._log(f"BEFORE сохранён: {path}"),self.status.set("BEFORE сохранён."),messagebox.showinfo("Готово",f"Baseline BEFORE сохранён:\n{path}")))
        self._thread(task)
    def run_after(self):
        def task():
            client,mapping,grouped,fid=self._inputs()
            before_path=Path(self.before_path.get()) if self.before_path.get() else None
            if not before_path or not before_path.exists():
                candidates=sorted(Path(self.output_dir.get()).glob("case_baseline_BEFORE_*.json"),key=lambda p:p.stat().st_mtime,reverse=True)
                if not candidates: raise ValueError("Сначала снимите baseline BEFORE.")
                before_path=candidates[0]
            before=json.loads(before_path.read_text(encoding="utf-8"))
            if str(before.get("fid")) != str(fid): raise ValueError(f"FID в BEFORE ({before.get('fid')}) не совпадает с текущим FID ({fid}).")
            after=client.snapshot(fid,grouped,self.fid_column.get())
            changes=compare_snapshots(before,after)
            comments=get_event_comments(self.xml_dir.get()) if self.xml_dir.get().strip() else []
            out=Path(self.output_dir.get())/f"case_checker_report_{now_stamp()}.html"
            html_report(before,after,changes,comments,str(out))
            after_path=Path(self.output_dir.get())/f"case_baseline_AFTER_{now_stamp()}.json"
            after_path.write_text(json.dumps(after,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
            self.after(0,lambda:(self._log(f"AFTER сохранён: {after_path}"),self._log(f"Отчёт создан: {out}"),self._log(f"Обнаружено изменений записей: {len(changes)}"),self.status.set("Сравнение завершено."),messagebox.showinfo("Готово",f"Отчёт создан:\n{out}\n\nИзменений записей: {len(changes)}")))
        self._thread(task)

if __name__ == "__main__":
    App().mainloop()
